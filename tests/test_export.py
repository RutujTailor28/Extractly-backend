"""Tests for Phase 10 export (CSV/XLSX/JSON/QuickBooks) — mocked data, no AI."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import date
from decimal import Decimal
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.document_service import DocumentServiceError
from app.services.export_service import (
    _csv_bytes,
    _json_payload,
    _xlsx_bytes,
    _zip_csv_bytes,
    generate_export,
    validate_date_range,
)
from app.services.quickbooks_mapping_service import map_to_quickbooks_rows


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def _invoice(**overrides):
    base = {
        "document_id": "doc-1",
        "filename": "inv.pdf",
        "vendor_name": "Acme Corp",
        "vendor_address": "1 Main St",
        "vendor_tax_id": "TAX-1",
        "invoice_number": "INV-1001",
        "invoice_date": "2026-09-01",
        "due_date": "2026-09-30",
        "po_number": "PO-9",
        "bill_to_name": None,
        "bill_to_address": None,
        "currency": "USD",
        "subtotal": "1000.00",
        "tax_amount": "100.00",
        "tax_rate": "10",
        "discount": "0",
        "shipping": "0",
        "total": "1100.00",
        "payment_terms": "Net 30",
        "bank_details": None,
        "status": "approved",
        "approved_at": "2026-09-03T00:00:00+00:00",
        "created_at": "2026-09-01T00:00:00+00:00",
    }
    base.update(overrides)
    return base


def _line(**overrides):
    base = {
        "document_id": "doc-1",
        "invoice_number": "INV-1001",
        "vendor_name": "Acme Corp",
        "line_number": 1,
        "description": 'Paper, "A4"',
        "quantity": "10",
        "unit": "box",
        "unit_price": "5.00",
        "line_tax": "5.00",
        "line_total": "55.00",
    }
    base.update(overrides)
    return base


def test_date_range_validation() -> None:
    validate_date_range(date(2026, 1, 1), date(2026, 2, 1))
    with pytest.raises(DocumentServiceError) as exc:
        validate_date_range(date(2026, 2, 1), date(2026, 1, 1))
    assert exc.value.code == "INVALID_DATE_RANGE"


def test_csv_escaping_and_unicode() -> None:
    from app.services.export_service import INVOICE_CSV_COLUMNS

    rows = [_invoice(vendor_name="Acmé, \"Corp\"", filename="réçu.pdf")]
    raw = _csv_bytes(rows, INVOICE_CSV_COLUMNS)
    text = raw.decode("utf-8-sig")
    assert "Acmé" in text
    reader = csv.DictReader(io.StringIO(text))
    parsed = list(reader)
    assert parsed[0]["vendor_name"] == 'Acmé, "Corp"'
    assert parsed[0]["total"] == "1100.00"


def test_csv_zip_with_line_items() -> None:
    content = _zip_csv_bytes([_invoice()], [_line()])
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        names = set(zf.namelist())
        assert names == {"invoices.csv", "line_items.csv"}
        invoices = zf.read("invoices.csv").decode("utf-8-sig")
        assert "INV-1001" in invoices


def test_xlsx_sheets_and_decimals() -> None:
    content = _xlsx_bytes([_invoice()], [_line()], include_line_items=True)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Invoices", "Line Items"]
    inv = wb["Invoices"]
    assert inv["A1"].value == "document_id"
    # total column
    headers = [c.value for c in inv[1]]
    total_idx = headers.index("total") + 1
    assert inv.cell(2, total_idx).value == pytest.approx(1100.0)


def test_json_structure_no_ai_metadata() -> None:
    payload = _json_payload(
        [_invoice()],
        [_line()],
        include_line_items=True,
        filters={"format": "json"},
    )
    assert payload["export_metadata"]["document_count"] == 1
    doc = payload["documents"][0]
    assert doc["invoice"]["amounts"]["total"] == pytest.approx(1100.0)
    blob = json.dumps(payload)
    assert "raw_response" not in blob
    assert "prompt" not in blob


def test_quickbooks_one_row_per_line_item() -> None:
    rows = map_to_quickbooks_rows(
        [_invoice()],
        [
            _line(line_number=1, description="Paper", line_total="50"),
            _line(line_number=2, description="Ink", line_total="40", quantity="2", unit_price="20"),
        ],
    )
    assert len(rows) == 2
    assert rows[0]["Vendor"] == "Acme Corp"
    assert rows[0]["Bill No"] == "INV-1001"
    assert rows[0]["Description"] == "Paper"
    assert rows[1]["Description"] == "Ink"


def test_quickbooks_no_line_items_still_one_row() -> None:
    rows = map_to_quickbooks_rows([_invoice()], [])
    assert len(rows) == 1
    assert rows[0]["Bill No"] == "INV-1001"
    assert rows[0]["Description"] == ""


def test_generate_export_rejects_unapproved_status() -> None:
    with pytest.raises(DocumentServiceError) as exc:
        generate_export(format="csv", status="needs_review")  # type: ignore[arg-type]
    assert exc.value.code == "DOCUMENT_NOT_APPROVED"


def test_export_api_empty_and_success(client: TestClient) -> None:
    with patch(
        "app.services.export_service.fetch_approved_documents",
        return_value=[],
    ), patch(
        "app.services.export_service._require_client",
        return_value=MagicMock(),
    ):
        response = client.get("/api/v1/export?format=csv")
    assert response.status_code == 404

    invoices = [_invoice()]
    lines = [_line()]
    with (
        patch(
            "app.services.export_service.fetch_approved_documents",
            return_value=[{"id": "doc-1", "status": "approved"}],
        ),
        patch(
            "app.services.export_service.build_export_rows",
            return_value=(invoices, lines),
        ),
        patch(
            "app.services.export_service._require_client",
            return_value=MagicMock(),
        ),
    ):
        csv_resp = client.get("/api/v1/export?format=csv&include_line_items=false")
        zip_resp = client.get("/api/v1/export?format=csv&include_line_items=true")
        xlsx_resp = client.get("/api/v1/export?format=xlsx")
        json_resp = client.get("/api/v1/export?format=json")

    assert csv_resp.status_code == 200
    assert "text/csv" in csv_resp.headers["content-type"]
    assert zip_resp.status_code == 200
    assert "zip" in zip_resp.headers["content-type"]
    assert xlsx_resp.status_code == 200
    assert json_resp.status_code == 200
    body = json_resp.json()
    assert body["export_metadata"]["document_count"] == 1


def test_quickbooks_preview_limited(client: TestClient) -> None:
    invoices = [_invoice(document_id=f"d{i}") for i in range(3)]
    lines = [_line(document_id=f"d{i}", line_number=j) for i in range(3) for j in range(40)]
    rows = map_to_quickbooks_rows(invoices, lines)
    assert len(rows) > 100

    with (
        patch(
            "app.services.quickbooks_mapping_service.build_quickbooks_dataset",
            return_value=(invoices, rows),
        ),
    ):
        response = client.get("/api/v1/export/quickbooks/preview")
    assert response.status_code == 200
    assert response.json()["row_count"] == 100


def test_export_limit_error(client: TestClient) -> None:
    with patch(
        "app.api.export.generate_export",
        side_effect=DocumentServiceError(
            "Export limit exceeded. Please narrow your filters.",
            status_code=422,
            code="EXPORT_LIMIT_EXCEEDED",
        ),
    ):
        response = client.get("/api/v1/export?format=json")
    assert response.status_code == 422
